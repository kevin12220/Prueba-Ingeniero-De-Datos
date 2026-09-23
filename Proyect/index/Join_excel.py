import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

production_orders = pd.read_csv(r"C:\Users\valen13\Downloads\Prueba Tecnica\Proyect\src\production_orders.csv")
shift = pd.read_csv(r"C:\Users\valen13\Downloads\Prueba Tecnica\Proyect\src\shifts.csv")

df = pd.merge(
    production_orders,
    shift,
    on="shift_id",
    how="inner"
)

df["unidades_optimas"] = (
    df["units_produced"] -
    df["units_defective"]
)

print(df[[
    "order_id",
    "machine_id",
    "shift_id",
    "supervisor",
    "units_produced",
    "units_defective",
    "unidades_optimas"
]])


resultado = df[
[
    "order_id",
    "machine_id",
    "shift_id",
    "supervisor",
    "units_produced",
    "units_defective",
    "unidades_optimas",
]
]
archivo_excel = r"C:\Users\valen13\Downloads\Prueba Tecnica\Proyect\src\unidades_optimas.xlsx"

resultado.to_excel(
    archivo_excel,
    index=False,
    sheet_name="Produccion"
)

wb = load_workbook(archivo_excel)
ws = wb["Produccion"]

for cell in ws[1]:
    cellont=Font(bold=True)
    
for column in ws.columns:
    max_length = 0
    column_letter = column[0].column_letter

    for cell in column:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    ws.column_dimensions[column_letter].width = max_length + 2
wb.save(archivo_excel)

print("Excel generado correctamente")